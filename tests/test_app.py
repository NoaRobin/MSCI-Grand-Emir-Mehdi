"""Tests hors connexion : aucun appel au vrai service MSCI."""
from __future__ import annotations

import io
import json
import threading
from http.server import ThreadingHTTPServer
from urllib.request import Request, urlopen

import pytest
from openpyxl import Workbook, load_workbook

from app import excel_io, mapping
from app import isin as isin_tools
from app.client import MSCIClient, MSCIHTTPError
from app.engine import Options, enrich_entity
from app.mock import MockClient


def report():
    return {"sources": {}, "catalog": {}, "refused_factors": {}, "warnings": []}


def rows(*codes):
    return [{"ISIN": c, "Client_ID": None, "Security_Status": None} for c in codes]


# ------------------------------------------------------------------ ISIN
def test_isin_validation():
    assert isin_tools.is_valid("GB00BTK05J60")
    assert isin_tools.is_valid("US0378331005")
    assert not isin_tools.is_valid("GB00BTK05J61")
    assert isin_tools.normalize(" gb00 btk05j60 ") == "GB00BTK05J60"
    assert "Cle" in isin_tools.problem("GB00BTK05J61")
    assert "caracteres" in isin_tools.problem("GB00")


# ------------------------------------------------------------------ mapping / template
def test_mapping_matches_template_exactly():
    columns = excel_io.template_columns()
    rules = mapping.load()
    assert list(rules["issuer"]) == columns["issuer"] and len(columns["issuer"]) == 94
    assert list(rules["fund"]) == columns["fund"] and len(columns["fund"]) == 123
    # Codes emetteur corriges : les scores de pilier ESG_*_SCORE sont des facteurs Fonds.
    assert rules["issuer"]["ESG_ENVIRONMENTAL_SCORE"]["candidates"][-1] == "ENVIRONMENTAL_PILLAR_SCORE"
    assert "IVA_RATING_DATE" in rules["issuer"]["ESG_RATING_DATE"]["candidates"]


# ------------------------------------------------------------------ lecture Excel
def _xlsx(sheets: dict[str, list[list]]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for name, data in sheets.items():
        ws = wb.create_sheet(name)
        for line in data:
            ws.append(line)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_read_team_input_format():
    content = _xlsx({"Titres": [["ISIN", "Client_ID"], ["GB00BTK05J60", "C1"], [None, None], ["ca3359341052", None]],
                     "Fonds": [["ISIN"], ["LU0274208692"]]})
    data = excel_io.read_isins("input.xlsx", content)
    assert [r["ISIN"] for r in data["issuer"]] == ["GB00BTK05J60", "CA3359341052"]
    assert data["issuer"][0]["Client_ID"] == "C1"
    assert [r["ISIN"] for r in data["fund"]] == ["LU0274208692"]


def test_read_headerless_column_and_offset_header():
    content = _xlsx({"Feuil1": [["GB00BTK05J60"], ["CA3359341052"]],
                     "Portefeuille": [["Rapport"], [], ["Nom", "Code ISIN"], ["Anglo", "GB00BTK05J60"]]})
    data = excel_io.read_isins("x.xlsx", content)
    assert [r["ISIN"] for r in data["auto"]] == ["GB00BTK05J60", "CA3359341052", "GB00BTK05J60"]


def test_read_csv_and_reject_unknown():
    data = excel_io.read_isins("l.csv", "ISIN;Client_ID\nGB00BTK05J60;A\n".encode())
    assert data["auto"][0]["Client_ID"] == "A"
    with pytest.raises(ValueError):
        excel_io.read_isins("l.pdf", b"x")


# ------------------------------------------------------------------ export
def test_export_keeps_template_and_forces_text():
    data = excel_io.export({"issuer": [{"ISSUER_NAME": "=HYPERLINK(\"x\")", "ESG_SCORE": 0, "UNGC_COMPLIANCE": False}],
                            "fund": []})
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Template Import Titres", "Template Import Fonds"]
    ws = wb["Template Import Titres"]
    assert [c.value for c in ws[1]] == excel_io.template_columns()["issuer"]
    assert ws["A2"].data_type == "s"
    headers = [c.value for c in ws[1]]
    assert ws.cell(2, headers.index("ESG_SCORE") + 1).value == 0
    assert ws.cell(2, headers.index("UNGC_COMPLIANCE") + 1).value is False


# ------------------------------------------------------------------ moteur
class FakeClient:
    """Client scripte : reponses par (version, endpoint)."""
    workers = 2

    def __init__(self, handlers, catalog=None):
        self.handlers, self.catalog, self.calls, self.audit = handlers, catalog, [], []

    def factors(self, entity, version="v4"):
        if self.catalog is None:
            raise MSCIHTTPError(403, "v4/metadata/factors")
        return {f: {} for f in self.catalog}

    def products(self, entity):
        return ["Fund Product"]

    def fetch_all(self, endpoint, payload, *, version="v4", extra_keys=()):
        self.calls.append((version, endpoint, payload))
        handler = self.handlers.get((version, endpoint))
        if handler is None:
            raise MSCIHTTPError(403, f"{version}{endpoint}", "refuse")
        return handler(payload)


def body(data, messages=(), unresolved=()):
    return [{"messages": list(messages), "result": {
        "response_metadata": {"unresolved_identifiers": list(unresolved)}, "data": data}}]


def test_cascade_fills_gaps_with_v3_then_history_and_keeps_zero():
    def v4(payload):
        ids = payload["issuer_identifier_list"]
        assert payload["inherit_missing_values"] is True
        return body([{"requested_id": i, "issuer_metadata": [{"ISSUER_NAME": "ANGLO", "ISSUERID": "IID1"}],
                      "factors": [{"name": "ESG_RATING_DATE", "value": None},
                                  {"name": "CARBON_GOVERNMENT_CO2_INTENSITY_GDP_TONPERMN", "value": 0}]}
                     for i in ids], messages=["The following factors were not recognized or not available "
                                              "to this user and were ignored: IVA_COMPANY_RATING, TRIR."])

    def v3(payload):
        assert "IVA_COMPANY_RATING" in payload["factor_name_list"]
        assert "CARBON_GOVERNMENT_CO2_INTENSITY_GDP_TONPERMN" not in payload["factor_name_list"]
        return body([{"requested_id": i, "factors": [{"name": "IVA_COMPANY_RATING", "value": "AA"}]}
                     for i in payload["issuer_identifier_list"]])

    def hist(payload):
        assert "start_date" in payload
        return body([{"requested_id": i, "factors": [{"name": "TRIR", "data_values": [
            {"value": 1.5, "as_of_date": "2026-09-30"}, {"value": 2.0, "as_of_date": "2026-10-28"}]}]}
            for i in payload["issuer_identifier_list"]])

    client = FakeClient({("v4", "/issuers"): v4, ("v3", "/issuers"): v3, ("v4", "/issuers/history"): hist},
                        catalog=["ISSUER_NAME", "ISSUER_TICKER", "ISSUERID", "ISSUER_ISIN", "ESG_RATING_DATE",
                                 "CARBON_GOVERNMENT_CO2_INTENSITY_GDP_TONPERMN", "IVA_COMPANY_RATING", "TRIR"])
    rep = report()
    out = enrich_entity(client, "issuer", rows("GB00BTK05J60", "GB00BTK05J61", "GB00BTK05J60"),
                        mapping.load()["issuer"], Options(), rep, lambda *a: None)
    assert len(out) == 3  # ordre et doublons conserves
    first = out[0]
    assert first["values"]["ISSUER_NAME"] == "ANGLO"
    assert first["values"]["CARBON_GOVERNMENT_CO2_INTENSITY_GDP_TONPERMN"] == 0
    assert first["values"]["ESG_RATING"] == "AA" and first["sources"]["ESG_RATING"] == "v3"
    assert first["values"]["TRIR"] == 2.0 and first["as_of"]["TRIR"] == "2026-10-28"
    assert first["state"] == "partial"
    assert out[1]["state"] == "invalid_isin"
    assert rep["refused_factors"]["issuer"]["v4"] == ["IVA_COMPANY_RATING", "TRIR"]
    # La V4 n'est interrogee que pour les codes presents a son catalogue
    v4_payload = next(p for v, e, p in client.calls if v == "v4" and e == "/issuers")
    assert "ENVIRONMENTAL_PILLAR_SCORE" not in v4_payload["factor_name_list"]


def test_refused_secondary_source_is_not_fatal():
    client = FakeClient({("v4", "/issuers"): lambda p: body(
        [{"requested_id": "GB00BTK05J60", "factors": [{"name": "ISSUER_NAME", "value": "A"}]}])}, catalog=None)
    rep = report()
    out = enrich_entity(client, "issuer", rows("GB00BTK05J60"), mapping.load()["issuer"], Options(), rep,
                        lambda *a: None)
    assert out[0]["state"] == "reference_only"
    assert rep["sources"]["issuer"]["v3"]["state"] == "refused"
    assert rep["warnings"]  # catalogue indisponible signale


def test_v4_server_error_stops_run():
    def boom(payload):
        raise MSCIHTTPError(500, "v4/issuers")
    with pytest.raises(MSCIHTTPError):
        enrich_entity(FakeClient({("v4", "/issuers"): boom}, catalog=[]), "issuer", rows("GB00BTK05J60"),
                      mapping.load()["issuer"], Options(), report(), lambda *a: None)


def test_fund_isin_resolved_through_coverage():
    def funds(payload):
        ids = payload["fund_identifier_list"]
        if ids == ["LU0274208692"]:
            return body([], unresolved=["LU0274208692"])
        assert ids == ["UIF000000000001"]
        return body([{"requested_id": "UIF000000000001", "factors": [{"name": "FUND_ESG_RATING", "value": "A"}]}])

    def coverage(payload):
        assert payload["entity_type"] == "fund"
        return body([{"MDSUID": "UIF000000000001", "ISIN": "LU0274208692", "FUND_NAME": "FUND X"}])

    client = FakeClient({("v4", "/funds"): funds, ("v4", "/coverage"): coverage},
                        catalog=["FUND_NAME", "FUND_ISIN", "FUND_ESG_RATING"])
    out = enrich_entity(client, "fund", rows("LU0274208692"), mapping.load()["fund"],
                        Options(use_v3=False, use_history=False), report(), lambda *a: None)
    values = out[0]["values"]
    assert out[0]["identifier"] == "UIF000000000001"
    assert values["FUND_ESG_RATING"] == "A" and values["FUND_NAME"] == "FUND X"
    assert values["FUND_ISIN"] == "LU0274208692"


def test_mock_full_run_produces_workbook():
    rules = mapping.load()
    out = enrich_entity(MockClient(0), "issuer", rows("GB00BTK05J60", "XS1556395710"), rules["issuer"],
                        Options(), report(), lambda *a: None)
    assert out[0]["esg_values"] > 0 and out[1]["state"] == "unresolved_identifier"
    data = excel_io.export({"issuer": [r["values"] for r in out]})
    assert load_workbook(io.BytesIO(data))["Template Import Titres"].max_row == 2


# ------------------------------------------------------------------ client HTTP
class FakeResponse:
    def __init__(self, status, payload, headers=None):
        self.status_code, self._payload, self.headers = status, payload, headers or {}

    def json(self):
        return self._payload


class FakeSession:
    def __init__(self, script):
        self.script, self.seen = script, []

    def request(self, method, url, **kwargs):
        self.seen.append((url, kwargs.get("json")))
        return self.script(url, kwargs.get("json"), len(self.seen))

    def close(self):
        pass


def test_client_token_reuse_401_refresh_and_parallel_batches():
    tokens = iter(["t1", "t2", "t3"])
    hits = {"401": False}

    def script(url, payload, n):
        if url.endswith("/oauth/token"):
            return FakeResponse(200, {"access_token": next(tokens), "expires_in": 86400})
        if not hits["401"]:
            hits["401"] = True
            return FakeResponse(401, {})
        batch = (payload or {}).get("batch_id", 1)
        return FakeResponse(200, {"code": 200, "messages": [], "result": {
            "response_metadata": {"total_number_of_batches": 3, "current_batch": batch,
                                  "data_request_id": "rid"}, "data": [{"batch": batch}]}})

    session = FakeSession(script)
    client = MSCIClient("id", "secret", session_factory=lambda: session, sleep=lambda s: None)
    pages = client.fetch_all("/issuers", {"issuer_identifier_list": ["X"]})
    assert [p["result"]["data"][0]["batch"] for p in pages] == [1, 2, 3]
    token_calls = [u for u, _ in session.seen if u.endswith("/oauth/token")]
    assert len(token_calls) == 2  # un renouvellement apres 401, puis reutilisation
    assert all("secret" not in json.dumps(entry) for entry in client.audit)


def test_client_retries_429_then_raises_typed_error():
    def script(url, payload, n):
        if url.endswith("/oauth/token"):
            return FakeResponse(200, {"access_token": "t", "expires_in": 100})
        return FakeResponse(429 if n < 3 else 403, {"messages": ["no subscription"]}, {"Retry-After": "0"})

    client = MSCIClient("id", "s", session_factory=lambda: FakeSession(script), sleep=lambda s: None)
    with pytest.raises(MSCIHTTPError) as err:
        client.call("POST", "/issuers", version="v3", json={})
    assert err.value.status == 403 and "v3/issuers" in str(err.value)


# ------------------------------------------------------------------ serveur
def test_server_end_to_end_demo(monkeypatch):
    from app import server
    monkeypatch.setenv("MSCI_DEMO", "1")
    server.Handler.state = server.State()
    httpd = ThreadingHTTPServer(("127.0.0.1", 0), server.Handler)
    threading.Thread(target=httpd.serve_forever, daemon=True).start()
    base = f"http://127.0.0.1:{httpd.server_address[1]}"
    try:
        status = json.load(urlopen(base + "/api/status"))
        assert status["demo"] is True and "secret" not in json.dumps(status)
        req = Request(base + "/api/run", data=json.dumps({"issuer": ["GB00BTK05J60"]}).encode(),
                      headers={"Content-Type": "application/json"})
        job = json.load(urlopen(req))["id"]
        for _ in range(100):
            data = json.load(urlopen(f"{base}/api/jobs/{job}"))
            if data["state"] != "running":
                break
            threading.Event().wait(0.1)
        assert data["state"] == "done", data
        workbook = load_workbook(io.BytesIO(urlopen(f"{base}/api/jobs/{job}/excel").read()))
        assert workbook.sheetnames == ["Template Import Titres", "Template Import Fonds"]
        forged = Request(base + "/api/demo", data=b"{}", headers={"Origin": "http://evil.example"})
        with pytest.raises(Exception):
            urlopen(forged)
    finally:
        httpd.shutdown()


def test_inherit_parameter_dropped_when_rejected():
    seen = []

    def v4(payload):
        seen.append(dict(payload))
        if "inherit_missing_values" in payload:
            raise MSCIHTTPError(400, "v4/issuers", "Invalid parameter")
        return body([{"requested_id": "GB00BTK05J60", "factors": [{"name": "ISSUER_NAME", "value": "A"}]}])

    rep = report()
    out = enrich_entity(FakeClient({("v4", "/issuers"): v4}, catalog=["ISSUER_NAME"]), "issuer",
                        rows("GB00BTK05J60"), mapping.load()["issuer"],
                        Options(use_v3=False, use_history=False), rep, lambda *a: None)
    assert out[0]["values"]["ISSUER_NAME"] == "A"
    assert "inherit_missing_values" not in seen[-1] and rep["warnings"]
